"""Slide template for the NeuS hands-on session, 7th EnzymeML Workshop 2026: black on white, 16:9,
Arial, with the University of Freiburg wordmark and the EnzymeML mark. Adapted from
~/Documents/presentations/enzymeml-workshop-2026/build_template.py (only logos, footer and output differ).

Run (from the repository root):  uv run --with python-pptx --with pillow --with lxml python slides/build_template.py
Writes slides/template.pptx: one master, eight layouts, no slides. build_deck.py fills it.
"""
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER as PH
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.oxml.shapes.picture import CT_Picture
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
LOGOS = HERE / "assets/logos"
OUT = HERE / "template.pptx"

INK, DIM, FONT = "0A0A0A", "6F6F6F", "Arial"
W, H, X = 13.333, 7.5, 0.9                        # slide and side margin, inches
CW = W - 2 * X
FOOTER = "7th EnzymeML Workshop · Rüdesheim 2026"
TABLE_STYLE = "{9D7B26C5-4107-4FEC-AEDC-1716B250EE53}"          # Light Style 1: hairline rules, no fills
# Freiburg wordmark (wide) first, then the EnzymeML mark (roughly square); heights relative to the mark
MARKS = [("uni-freiburg-wordmark-black.png", 0.38), ("enzymeml-logo.png", 1.0)]
E = lambda v: int(Inches(v))


# ---- xml helpers --------------------------------------------------------------------------
def lvl(n, size, color=INK, bold=False, algn="l", bullet=False, caps=False, spc=0, line=None, before=0):
    """One <a:lvlNpPr> of a list style."""
    mar = E(0.3 * n) if bullet else 0
    ind = -E(0.3) if bullet else 0
    ln = f'<a:lnSpc><a:spcPct val="{line * 1000}"/></a:lnSpc>' if line else ""
    bu = '<a:buFont typeface="Arial"/><a:buChar char="•"/>' if bullet else "<a:buNone/>"
    cap = ' cap="all"' if caps else ""
    return (f'<a:lvl{n}pPr marL="{mar}" indent="{ind}" algn="{algn}">{ln}<a:spcBef><a:spcPts val="{before * 100}"/></a:spcBef>{bu}'
            f'<a:defRPr sz="{size * 100}" b="{int(bold)}" spc="{spc}"{cap}><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="{FONT}"/></a:defRPr></a:lvl{n}pPr>')


def style(ph, *levels, anchor="t"):
    """Replace a placeholder's list style (empty = inherit the master) and zero its insets."""
    body = ph._element.txBody
    bodyPr = body.find(qn("a:bodyPr"))
    for k, v in dict(lIns="0", tIns="0", rIns="0", bIns="0", anchor=anchor).items():
        bodyPr.set(k, v)
    old = body.find(qn("a:lstStyle"))
    old.addprevious(parse_xml(f'<a:lstStyle {nsdecls("a")}>{"".join(levels)}</a:lstStyle>'))
    body.remove(old)


def place(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = E(x), E(y), E(w), E(h)


def next_id(tree):
    return max(int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))) + 1


def picture(part, tree, path, x, y, w, h):
    _, rId = part.get_or_add_image_part(str(path))
    name = f"Logo {path.stem.replace('-black', '')}"
    tree.append(CT_Picture.new_pic(next_id(tree), name, name, rId, E(x), E(y), E(w), E(h)))


def text(tree, x, y, w, h, runs, algn="l", name="Footer text"):
    tree.append(parse_xml(
        f'<p:sp {nsdecls("a", "p")}><p:nvSpPr><p:cNvPr id="{next_id(tree)}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(w)}" cy="{E(h)}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="{algn}"/>{runs}</a:p></p:txBody></p:sp>'))


def rpr(size, color, bold=False, spc=0):
    return f'<a:rPr lang="en-GB" sz="{size * 100}" b="{int(bold)}" spc="{spc}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:latin typeface="{FONT}"/></a:rPr>'


def text_ph(layout, idx, name, y, h, anchor, level):
    """A full-width body placeholder with its own list style."""
    tree = layout.shapes._spTree
    tree.append(parse_xml(
        f'<p:sp {nsdecls("a", "p")}><p:nvSpPr><p:cNvPr id="{next_id(tree)}" name="{name}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        f'<p:nvPr><p:ph type="body" sz="quarter" idx="{idx}"/></p:nvPr></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{E(X)}" y="{E(y)}"/><a:ext cx="{E(CW)}" cy="{E(h)}"/></a:xfrm></p:spPr>'
        f'<p:txBody><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anchor}"/><a:lstStyle>{level}</a:lstStyle>'
        f'<a:p><a:r><a:rPr lang="en-GB"/><a:t>{name}</a:t></a:r></a:p></p:txBody></p:sp>'))


KICKER, SUBTITLE, CITATION = 13, 15, 16          # placeholder idx, shared with build_deck.py
kicker = lambda lay, y: text_ph(lay, KICKER, "Kicker", y, 0.3, "b", lvl(1, 11, DIM, bold=True, caps=True, spc=300))
subtitle = lambda lay: text_ph(lay, SUBTITLE, "Subtitle", 1.5, 0.6, "t", lvl(1, 18, DIM, line=100))
# references: small, just above the footer, growing upwards
citation = lambda lay: text_ph(lay, CITATION, "Citation", 6.35, 0.42, "b", lvl(1, 9, DIM, line=100))


def logos(part, tree, left, mid, seal, center=False):
    """The two marks in a row, vertically centred on `mid`; returns the right edge."""
    sizes = [(f, seal * r * Image.open(LOGOS / f).width / Image.open(LOGOS / f).height, seal * r) for f, r in MARKS]
    gaps = [seal * 0.55, 0]
    x = left - (sum(w for _, w, _ in sizes) + sum(gaps)) / 2 if center else left
    for (f, w, h), gap in zip(sizes, gaps):
        picture(part, tree, LOGOS / f, x, mid - h / 2, w, h)
        x += w + gap
    return x


# ---- theme and master ---------------------------------------------------------------------
def theme(prs):
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    scheme = root.find(f".//{a}clrScheme")
    scheme.set("name", "Black and white")
    for name, val in dict(dk1=INK, lt1="FFFFFF", dk2=INK, lt2="F2F2F2", accent1=INK, accent2=DIM, accent3="A6A6A6",
                          accent4="404040", accent5="8C8C8C", accent6="D9D9D9", hlink=INK, folHlink=DIM).items():
        el = scheme.find(a + name)
        el.clear()
        etree.SubElement(el, a + "srgbClr", val=val)
    for font in root.iter(f"{a}majorFont", f"{a}minorFont"):
        font.find(a + "latin").set("typeface", FONT)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def strip_footers(owner):
    for ph in list(owner.placeholders):
        if ph.element.ph_type in (PH.DATE, PH.FOOTER, PH.SLIDE_NUMBER):
            ph.element.getparent().remove(ph.element)


def master(prs):
    m = prs.slide_master
    tx = m._element.find(qn("p:txStyles"))
    tx.replace(tx.find(qn("p:titleStyle")),
               parse_xml(f'<p:titleStyle {nsdecls("a", "p")}>{lvl(1, 34, bold=True, spc=-60, line=92)}</p:titleStyle>'))
    tx.replace(tx.find(qn("p:bodyStyle")), parse_xml(
        f'<p:bodyStyle {nsdecls("a", "p")}>{lvl(1, 20, bullet=True, line=105, before=10)}'
        f'{lvl(2, 16, DIM, bullet=True, before=6)}{lvl(3, 14, DIM, bullet=True, before=4)}</p:bodyStyle>'))
    strip_footers(m)
    for ph in m.placeholders:
        if ph.element.ph_type == PH.TITLE:
            place(ph, X, 0.9, CW, 0.9); style(ph)
        else:
            place(ph, X, 2.2, CW, 4.1); style(ph)
    tree = m.shapes._spTree
    right = logos(m.part, tree, X, 7.05, 0.28)
    text(tree, right + 0.3, 6.9, 5.0, 0.3, f"<a:r>{rpr(9, DIM, spc=40)}<a:t>{FOOTER}</a:t></a:r>")
    text(tree, W - X - 1.0, 6.9, 1.0, 0.3,
         f'<a:fld id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum">{rpr(9, INK, bold=True)}<a:t>‹#›</a:t></a:fld>',
         algn="r", name="Slide number")


# ---- layouts -------------------------------------------------------------------------------
def layouts(prs):
    L = list(prs.slide_layouts)                      # the default template's eleven, by index
    for i in (7, 9, 10):                             # content with caption, both vertical-text layouts
        prs.slide_layouts.remove(L[i])
    for lay in L:
        strip_footers(lay)
    ph = lambda lay, idx: next(p for p in lay.placeholders if p.placeholder_format.idx == idx)
    def std_title(lay):                              # kicker, one-line title, subtitle, citation
        place(ph(lay, 0), X, 0.9, CW, 0.6); style(ph(lay, 0))
        kicker(lay, 0.55); subtitle(lay); citation(lay)
    names = {0: "Cover", 1: "Content", 2: "Section", 3: "Two columns", 4: "End", 5: "Title only", 6: "Blank", 8: "Figure"}
    for i, n in names.items():
        L[i]._element.cSld.set("name", n)

    cover = L[0]; cover._element.set("showMasterSp", "0")
    place(ph(cover, 0), X, 2.75, CW * 0.88, 2.3); style(ph(cover, 0), lvl(1, 46, bold=True, spc=-80, line=92))
    place(ph(cover, 1), X, 5.2, CW, 1.3)
    style(ph(cover, 1), lvl(1, 20, line=110), lvl(2, 12, DIM, before=6))
    kicker(cover, 2.3)
    logos(cover.part, cover.shapes._spTree, X, 1.2, 0.95)

    content = L[1]; std_title(content)
    place(ph(content, 1), X, 2.2, CW, 4.1); style(ph(content, 1))

    section = L[2]
    place(ph(section, 0), X, 2.8, CW, 1.6); style(ph(section, 0), lvl(1, 60, bold=True, spc=-100, line=90))
    place(ph(section, 1), X, 4.0, CW, 1.0); style(ph(section, 1), lvl(1, 22, DIM))
    kicker(section, 2.45)

    two = L[3]; std_title(two)
    place(ph(two, 1), X, 2.2, 5.5, 4.1); style(ph(two, 1))
    place(ph(two, 2), W - X - 5.5, 2.2, 5.5, 4.1); style(ph(two, 2))

    end = L[4]; end._element.set("showMasterSp", "0")
    for p in list(end.placeholders):
        if p.placeholder_format.idx not in (0, 1):
            p.element.getparent().remove(p.element)
    place(ph(end, 0), X, 2.2, CW, 1.3); style(ph(end, 0), lvl(1, 60, bold=True, spc=-100, algn="ctr"), anchor="b")
    place(ph(end, 1), X, 3.7, CW, 1.2)
    style(ph(end, 1), lvl(1, 22, DIM, algn="ctr"), lvl(2, 11, INK, bold=True, caps=True, spc=300, algn="ctr", before=14))
    logos(end.part, end.shapes._spTree, W / 2, 5.95, 0.7, center=True)

    std_title(L[5])
    citation(L[6])
    fig = L[8]; std_title(fig)
    place(ph(fig, 1), X, 2.2, CW, 3.5)
    place(ph(fig, 2), X, 5.8, CW, 0.5); style(ph(fig, 2), lvl(1, 14, bullet=True, before=2))
    # order in the layout gallery
    ids = prs.slide_master._element.find(qn("p:sldLayoutIdLst"))
    by_rid = {prs.slide_master.part.related_part(e.get(qn("r:id"))).slide_layout.name: e for e in ids}
    for name in ["Cover", "Section", "Content", "Figure", "Two columns", "Title only", "Blank", "End"]:
        ids.append(by_rid[name])


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W), E(H)
    prs.core_properties.title = "From measurement to kinetic modelling"
    theme(prs)
    styles = prs.part.part_related_by(RT.TABLE_STYLES)            # default table style for new tables
    styles._blob = styles.blob.replace(b"{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}", TABLE_STYLE.encode())
    master(prs)
    layouts(prs)
    prs.save(OUT)
    print(f"wrote {OUT} ({len(prs.slides)} slides, layouts: {', '.join(l.name for l in prs.slide_layouts)})")


if __name__ == "__main__":
    main()
